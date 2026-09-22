"""
app/main.py
FastAPI web application for Face Recognition Attendance System.
Provides endpoints for webcam frame recognition, employee enrollment, active presence monitoring,
health checks (for Render keep-alive pings), and static frontend serving.
"""

import os
import sys
import io
import datetime
import time
import psycopg2
import traceback
from pathlib import Path
from typing import Optional, List

from fastapi import FastAPI, HTTPException, Header, Depends, status, Request
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

import app.database as database
from recognizer import (
    FaceRecognizer,
    CONFIDENCE_FLOOR,
    decode_base64_image,
    compute_encoding_from_bgr,
    compute_encoding_and_box,
    compute_averaged_encoding_from_images,
    validate_enrollment_sample,
)


ADMIN_TOKEN = os.environ.get("ADMIN_TOKEN", "admin123")
if os.environ.get("ADMIN_TOKEN") is None:
    print("[WARNING] ADMIN_TOKEN environment variable is not set. Using default fallback token: 'admin123'")

KIOSK_SECRET = os.environ.get("KIOSK_SECRET", "facepulse_kiosk_client_default")
if os.environ.get("KIOSK_SECRET") is None:
    print("[WARNING] KIOSK_SECRET environment variable is not set. Using default fallback: 'facepulse_kiosk_client_default'")


app = FastAPI(
    title="Face Pulse Attendance API",
    description="Live face recognition attendance web service",
    version="1.0.0",
)

@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception):
    err_tb = traceback.format_exc()
    print(f"[Unhandled Server Error] {request.method} {request.url.path}: {err_tb}")
    return JSONResponse(
        status_code=500,
        content={"detail": f"Server error: {str(exc)}", "type": type(exc).__name__}
    )

# Enable CORS for web client access
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Global face recognizer instance
recognizer_instance: Optional[FaceRecognizer] = None

# In-memory throttles to prevent synchronous database write contention on high-frequency frame streams
_last_logged_presence = {}  # employee_id -> epoch timestamp
_presence_cache = []
_presence_cache_time = 0.0

# Change 2: Consecutive recognition checks required per active track ID before logging to presence_log
REQUIRED_CONSECUTIVE_MATCHES = 3
_track_history = {}  # track_id -> {"history": list of employee_ids, "last_seen": float}



@app.on_event("startup")
def startup_event():
    database.init_db()
    get_recognizer()
    print("Database initialized and face recognizer loaded successfully.")


def get_recognizer() -> FaceRecognizer:
    global recognizer_instance
    if recognizer_instance is None:
        database.init_db()
        recognizer_instance = FaceRecognizer()
    return recognizer_instance


# Request models
class RecognizeRequest(BaseModel):
    image: str  # Base64 data URL string


class EnrollRequest(BaseModel):
    name: str
    image: Optional[str] = None  # Single Base64 data URL string
    images: Optional[List[str]] = None  # List of 15 Base64 data URL strings for multi-sample guided enrollment
    admin_token: str


class VerifySampleRequest(BaseModel):
    image: str  # Base64 data URL string
    existing_images: Optional[List[str]] = None  # Already collected sample base64 strings for duplicate check


class VerifyTokenRequest(BaseModel):
    admin_token: str


# Helper dependency to check admin token
def verify_admin_token(x_admin_token: Optional[str] = Header(None), token: Optional[str] = None):
    provided = x_admin_token or token
    if not provided or provided != ADMIN_TOKEN:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid or missing admin token.",
        )
    return True


# ---------------- API Endpoints ----------------

@app.get("/health")
def health_check():
    """
    Keep-alive endpoint for ping services (e.g. cron-job.org or UptimeRobot)
    to prevent Render free tier instance spin-down.
    """
    return {
        "status": "ok",
        "service": "face-pulse-attendance",
        "timestamp": datetime.datetime.now().isoformat(),
    }


@app.post("/api/verify-token")
def verify_token(req: VerifyTokenRequest):
    """Verify if the supplied admin token is valid."""
    is_valid = req.admin_token == ADMIN_TOKEN
    return {"valid": is_valid}


@app.post("/api/recognize")
def recognize_frame(
    req: RecognizeRequest,
    x_kiosk_client: Optional[str] = Header(None),
):
    """
    Accepts a base64 camera frame, performs fast face detection + IoU/centroid tracking + recognition,
    logs presence in DB with in-memory throttling, and returns tracked bounding boxes and names.
    Protected by lightweight X-Kiosk-Client header matching KIOSK_SECRET.
    """
    if not x_kiosk_client or x_kiosk_client != KIOSK_SECRET:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Forbidden: Invalid or missing X-Kiosk-Client header.",
        )

    recognizer = get_recognizer()
    t_start = time.perf_counter()

    try:
        frame_bgr = decode_base64_image(req.image)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Invalid base64 image data: {str(e)}")

    result = recognizer.process_bgr_frame(frame_bgr)

    # In-memory throttle and Change 2 consecutive-frame agreement tracking
    global _presence_cache, _presence_cache_time, _track_history
    now_ts = time.time()

    # Clean up stale tracks from memory if dictionary grows large
    if len(_track_history) > 100:
        stale_tids = [tid for tid, data in _track_history.items() if now_ts - data.get("last_seen", 0.0) > 60.0]
        for tid in stale_tids:
            del _track_history[tid]

    for m in result.get("matches", []):
        track_id = m.get("track_id")
        emp_id = m.get("employee_id")
        name = m.get("name")
        conf = m.get("confidence")

        # Determine if this frame produced an accepted identity match (passed Change 1 checks)
        is_accepted_match = (
            emp_id is not None and
            name != "Unknown" and
            (conf is None or conf >= CONFIDENCE_FLOOR)
        )
        current_identity = emp_id if is_accepted_match else None

        # Track consecutive recognition results per track ID
        if track_id is not None:
            if track_id not in _track_history:
                _track_history[track_id] = {"history": [], "last_seen": now_ts}

            track_record = _track_history[track_id]
            track_record["last_seen"] = now_ts
            history = track_record["history"]

            if current_identity is None:
                # Identity rejected (failed margin check, distance check, or unknown) -> reset streak
                history.clear()
            else:
                if history and history[-1] != current_identity:
                    # Identity flipped mid-sequence (e.g. matched as Alice then Bob) -> reset counter
                    history.clear()
                history.append(current_identity)
                if len(history) > REQUIRED_CONSECUTIVE_MATCHES:
                    history.pop(0)

            # Only write to presence_log when the SAME identity has been accepted for N consecutive checks
            is_confirmed = (
                len(history) >= REQUIRED_CONSECUTIVE_MATCHES and
                all(x == current_identity for x in history) and
                current_identity is not None
            )

            if is_confirmed:
                if now_ts - _last_logged_presence.get(current_identity, 0.0) >= 30.0:
                    _last_logged_presence[current_identity] = now_ts
                    try:
                        database.log_presence(current_identity, distance=m.get("distance"), min_interval_seconds=30)
                        print(
                            f"[Attendance Confirmed] Track #{track_id} confirmed as '{name}' (ID={current_identity}) "
                            f"across {REQUIRED_CONSECUTIVE_MATCHES} consecutive checks. Logged to DB."
                        )
                    except Exception as e:
                        print(f"Error logging presence: {e}")

    # Cache get_recently_present for 2.0s to avoid expensive DB queries on every video frame
    if now_ts - _presence_cache_time >= 2.0:
        try:
            _presence_cache = database.get_recently_present(window_seconds=20)
            _presence_cache_time = now_ts
        except Exception as e:
            print(f"Error refreshing presence cache: {e}")

    result["present"] = _presence_cache
    result["proc_ms"] = round((time.perf_counter() - t_start) * 1000, 1)
    return result


@app.post("/api/enroll-check")
def verify_enrollment_sample(req: VerifySampleRequest):
    """
    Verifies if a candidate camera frame contains exactly 1 detectable face
    and rejects duplicate/near-identical pose frames compared to existing samples.
    """
    try:
        frame_bgr = decode_base64_image(req.image)
        valid, reason, box = validate_enrollment_sample(frame_bgr, req.existing_images)
        return {"valid": valid, "reason": reason, "box": box}
    except Exception as e:
        return {"valid": False, "reason": f"Sample verification failed: {str(e)}"}


@app.post("/api/enroll")
def enroll_employee(req: EnrollRequest):
    """
    Enrolls a new person into the system using single or multi-sample guided enrollment.
    Requires valid admin_token. Computes averaged 128-d face encoding from samples.
    """
    print(f"[Enroll] Request received for name='{req.name}', admin_token_provided={'Yes' if req.admin_token else 'No'}")

    if req.admin_token != ADMIN_TOKEN:
        print(f"[Enroll REJECTED] Admin token mismatch. Received '{req.admin_token}', expected '{ADMIN_TOKEN}'")
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Unauthorized: Incorrect admin token. Please enter valid admin token.",
        )

    clean_name = req.name.strip()
    if not clean_name:
        raise HTTPException(status_code=400, detail="Name cannot be empty.")

    encoding = None
    samples_used = 1

    if req.images and len(req.images) > 0:
        print(f"[Enroll] Processing multi-sample enrollment for '{clean_name}' with {len(req.images)} images...")
        avg_enc, count = compute_averaged_encoding_from_images(req.images)
        if avg_enc is None or count == 0:
            print(f"[Enroll ERROR] Could not detect faces in multi-sample set for '{clean_name}'.")
            raise HTTPException(
                status_code=400,
                detail="Could not detect faces in the provided sample set. Please ensure good lighting and re-enroll.",
            )
        encoding = avg_enc
        samples_used = count
    elif req.image:
        print(f"[Enroll] Processing single-sample enrollment for '{clean_name}'...")
        try:
            frame_bgr = decode_base64_image(req.image)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Invalid image format: {str(e)}")

        encoding = compute_encoding_from_bgr(frame_bgr)
        if encoding is None:
            print(f"[Enroll ERROR] No face detected in single image for '{clean_name}'.")
            raise HTTPException(
                status_code=400,
                detail="No face detected in the image. Please position your face clearly in front of the camera and try again.",
            )
    else:
        raise HTTPException(status_code=400, detail="No image data provided for enrollment.")

    try:
        emp_id = database.add_employee(clean_name, encoding)
        print(f"[Enroll DB SUCCESS] Saved employee '{clean_name}' with ID={emp_id} to database.")
    except psycopg2.IntegrityError as e:
        print(f"[Enroll DB INTEGRITY ERROR] Duplicate employee name '{clean_name}': {e}")
        raise HTTPException(
            status_code=400,
            detail=f"Employee '{clean_name}' already exists in the database. Please choose a different name.",
        )
    except psycopg2.OperationalError as e:
        print(f"[Enroll DB OPERATIONAL ERROR] Database error saving '{clean_name}': {e}")
        raise HTTPException(
            status_code=503,
            detail=f"Database write failed due to connection/concurrency error: {str(e)}",
        )
    except Exception as e:
        print(f"[Enroll DB ERROR] Failed to save '{clean_name}': {e}")
        raise HTTPException(
            status_code=400,
            detail=f"Could not enroll user: {str(e)}",
        )

    # Refresh recognizer with the new employee encoding (non-fatal if reload fails)
    try:
        if recognizer_instance:
            recognizer_instance.refresh_known_faces()
            print(f"[Enroll RECOGNIZER SUCCESS] Refreshed recognizer with new employee '{clean_name}'.")
    except Exception as e:
        print(f"[Enroll RECOGNIZER WARNING] Failed to refresh known faces cache for '{clean_name}': {e}")

    return {
        "status": "success",
        "message": f"Enrolled '{clean_name}' successfully using {samples_used} face samples!",
        "employee_id": emp_id,
        "name": clean_name,
        "samples_used": samples_used,
    }



@app.get("/api/present")
def get_currently_present(window_seconds: int = 20):
    """Get list of employees recognized within the last `window_seconds`."""
    present = database.get_recently_present(window_seconds=window_seconds)
    return {"present": present}


def check_admin_token(
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = None,
):
    """Verifies that a valid ADMIN_TOKEN is provided via query parameter or header."""
    provided = admin_token or token or x_admin_token
    if not provided or provided != ADMIN_TOKEN:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Unauthorized: Invalid or missing admin token.",
        )


@app.get("/api/employees")
def get_employees(
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = Header(None),
):
    """Get list of all enrolled employees. Gated by ADMIN_TOKEN."""
    check_admin_token(admin_token, token, x_admin_token)
    employees = database.get_all_people()
    return {"employees": employees}


@app.delete("/api/employees/{employee_id}")
def delete_employee(
    employee_id: int,
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = Header(None),
):
    """Delete an enrolled employee by ID. Gated by ADMIN_TOKEN."""
    check_admin_token(admin_token, token, x_admin_token)
    database.delete_employee(employee_id)
    if recognizer_instance:
        recognizer_instance.refresh_known_faces()

    return {"status": "success", "message": "Employee deleted."}


@app.get("/api/activity-log")
def get_activity_log(
    limit: int = 100,
    name: Optional[str] = None,
    date: Optional[str] = None,
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = Header(None),
):
    """Get detailed timestamped activity log records with ML match distance & confidence scores. Gated by ADMIN_TOKEN."""
    check_admin_token(admin_token, token, x_admin_token)
    logs = database.get_activity_log(limit=limit, name_filter=name, date_filter=date)
    return {"logs": logs}


@app.get("/api/today-summary")
def get_today_summary(
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = Header(None),
):
    """Get attendance log summary for today. Gated by ADMIN_TOKEN."""
    check_admin_token(admin_token, token, x_admin_token)
    summary = database.get_today_summary()
    return {"summary": summary}


@app.get("/api/admin/status")
def get_admin_status(
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = Header(None),
):
    """System and database pool health check. Gated by ADMIN_TOKEN."""
    check_admin_token(admin_token, token, x_admin_token)
    pool_status = database.get_pool_status()
    employees = database.get_all_people()
    return {
        "status": "online",
        "pool": pool_status,
        "enrolled_count": len(employees),
    }


def format_time_clean(iso_ts: Optional[str]) -> str:
    """Helper to convert ISO timestamp to clean readable local/formatted time."""
    if not iso_ts:
        return "N/A"
    try:
        clean_str = str(iso_ts)
        if not clean_str.endswith("Z") and "+" not in clean_str and "-" not in clean_str[10:]:
            clean_str += "+00:00"
        dt = datetime.datetime.fromisoformat(clean_str)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(iso_ts)


@app.get("/api/report/today.xlsx")
def download_today_report_excel(
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = Header(None)
):
    """
    Downloads today's attendance summary report as a formatted Excel spreadsheet (.xlsx).
    Gated behind ADMIN_TOKEN authentication. Generated completely in memory.
    """
    check_admin_token(admin_token, token, x_admin_token)

    summary = database.get_today_summary()
    today_str = datetime.date.today().isoformat()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Today's Attendance"

    # Title Row
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = f"Attendance Report — {today_str}"
    title_cell.font = Font(size=14, bold=True, color="1E293B")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Header Row
    headers = ["Name", "First Seen", "Last Seen", "Total Detections"]
    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 24

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for col_idx in range(1, 5):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align if col_idx > 1 else left_align
        cell.border = thin_border

    # Data Rows
    row_num = 4
    if summary:
        for s in summary:
            f_seen = format_time_clean(s.get("first_seen"))
            l_seen = format_time_clean(s.get("last_seen"))
            dets = s.get("detections", 0)
            ws.append([s.get("name", "Unknown"), f_seen, l_seen, dets])

            row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if row_num % 2 == 0 else None
            for col_idx in range(1, 5):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.border = thin_border
                if row_fill:
                    cell.fill = row_fill
                cell.alignment = center_align if col_idx > 1 else left_align
            row_num += 1
    else:
        ws.merge_cells("A4:D4")
        empty_cell = ws["A4"]
        empty_cell.value = "No attendance records logged yet today."
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        empty_cell.font = Font(italic=True, color="64748B")

    # Column Widths
    col_widths = {"A": 24, "B": 24, "C": 24, "D": 18}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"attendance_report_{today_str}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/report/today.pdf")
def download_today_report_pdf(
    admin_token: Optional[str] = None,
    token: Optional[str] = None,
    x_admin_token: Optional[str] = Header(None)
):
    """
    Downloads today's attendance summary report as a PDF document.
    Gated behind ADMIN_TOKEN authentication. Generated completely in memory.
    """
    check_admin_token(admin_token, token, x_admin_token)

    summary = database.get_today_summary()
    today_str = datetime.date.today().isoformat()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1e293b"),
        alignment=1,  # Center
        spaceAfter=15,
    )

    elements = [
        Paragraph(f"Attendance Report — {today_str}", title_style),
        Spacer(1, 10),
    ]

    table_data = [["Name", "First Seen", "Last Seen", "Total Detections"]]
    if summary:
        for s in summary:
            f_seen = format_time_clean(s.get("first_seen"))
            l_seen = format_time_clean(s.get("last_seen"))
            dets = str(s.get("detections", 0))
            table_data.append([s.get("name", "Unknown"), f_seen, l_seen, dets])
    else:
        table_data.append(["No attendees logged today", "-", "-", "0"])

    col_widths = [140, 150, 150, 100]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("ALIGN", (0, 1), (0, -1), "LEFT"),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("TOPPADDING", (0, 1), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
    ]))

    elements.append(t)
    doc.build(elements)
    buf.seek(0)

    filename = f"attendance_report_{today_str}.pdf"
    return Response(
        content=buf.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )



# ---------------- Static Files & UI Mounting ----------------

BASE_DIR = Path(__file__).parent.parent
STATIC_DIR = BASE_DIR / "static"

if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


@app.get("/")
def read_root():
    index_file = STATIC_DIR / "index.html"
    if index_file.exists():
        return FileResponse(str(index_file))
    return {"message": "Face Pulse API is running. Static frontend not found."}


@app.get("/admin")
def read_admin():
    admin_file = STATIC_DIR / "admin.html"
    if admin_file.exists():
        return FileResponse(str(admin_file))
    return {"message": "Admin dashboard not found."}
